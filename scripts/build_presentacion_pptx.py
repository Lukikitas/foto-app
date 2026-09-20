#!/usr/bin/env python3
"""Build an editable widescreen PowerPoint of the Delivery La Plata deck."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "public" / "presentacion"
OUT = ASSETS / "Delivery-La-Plata.pptx"

BG = RGBColor(0x11, 0x11, 0x13)
CARD = RGBColor(0x1A, 0x1B, 0x1E)
LINE = RGBColor(0x2C, 0x2E, 0x33)
RED = RGBColor(0xE4, 0x00, 0x2B)
HEAD = RGBColor(0xF4, 0xF5, 0xF7)
TEXT = RGBColor(0xC3, 0xC6, 0xCE)
MUTED = RGBColor(0x8A, 0x8F, 0x99)
OK = RGBColor(0x3D, 0xCC, 0x7A)
WARN = RGBColor(0xF0, 0xB4, 0x29)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W = Inches(13.333)
H = Inches(7.5)
MARGIN = Inches(0.48)


def set_run(run, text, size=18, bold=False, color=TEXT, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", font)
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)


def write_paragraph(p, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, space_after=6):
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return run


def textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def card(slide, x, y, w, h, fill=CARD, line=LINE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    fill_shape(shape, fill)
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    # tighter corners
    try:
        shape.adjustments[0] = 0.06
    except Exception:
        pass
    return shape


def kicker(slide, text, y=Inches(0.28)):
    box, tf = textbox(slide, MARGIN, y, Inches(12.3), Inches(0.32))
    write_paragraph(tf.paragraphs[0], text.upper(), 11, True, RED, space_after=0)
    return box


def title(slide, text, y=Inches(0.52), size=28, w=Inches(12.3)):
    box, tf = textbox(slide, MARGIN, y, w, Inches(0.7))
    write_paragraph(tf.paragraphs[0], text, size, True, HEAD, space_after=0)
    return box


def lead(slide, text, y=Inches(1.18), w=Inches(12.3), size=16):
    box, tf = textbox(slide, MARGIN, y, w, Inches(0.7))
    write_paragraph(tf.paragraphs[0], text, size, False, MUTED, space_after=0)
    return box


def bullets(tf, items, size=15, color=TEXT, level0_size=None):
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 0
        p.space_after = Pt(8)
        p.space_before = Pt(0)
        if isinstance(item, tuple):
            bold_part, rest = item
            run = p.add_run()
            set_run(run, "•  ", size, False, RED)
            run = p.add_run()
            set_run(run, bold_part, size, True, HEAD)
            run = p.add_run()
            set_run(run, rest, size, False, color)
        else:
            run = p.add_run()
            set_run(run, "•  ", size, False, RED)
            run = p.add_run()
            set_run(run, item, size, False, color)


def picture(slide, path, x, y, w, h):
    return slide.shapes.add_picture(str(path), x, y, w, h)


def footer(slide, left="KFC · Delivery La Plata", right="delivery.star-app.com.ar"):
    box, tf = textbox(slide, MARGIN, Inches(7.18), Inches(6.5), Inches(0.24))
    write_paragraph(tf.paragraphs[0], left, 10, False, MUTED, space_after=0)
    box, tf = textbox(slide, Inches(7.2), Inches(7.18), Inches(5.6), Inches(0.24))
    write_paragraph(tf.paragraphs[0], right, 10, False, MUTED, PP_ALIGN.RIGHT, 0)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill_shape(bg, BG)
    # send background back
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return slide


def stat_card(slide, x, y, w, h, value, label, value_color=HEAD):
    card(slide, x, y, w, h)
    box, tf = textbox(slide, x + Inches(0.18), y + Inches(0.16), w - Inches(0.36), Inches(0.5))
    write_paragraph(tf.paragraphs[0], value, 26, True, value_color, space_after=0)
    box, tf = textbox(slide, x + Inches(0.18), y + Inches(0.66), w - Inches(0.36), Inches(0.5))
    write_paragraph(tf.paragraphs[0], label, 12, False, MUTED, space_after=0)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Cover
    s = blank_slide(prs)
    kicker(s, "KFC · Dirección · Lucas Rodriguez")
    title(s, "Delivery La Plata", y=Inches(0.55), size=40, w=Inches(6.4))
    box, tf = textbox(s, MARGIN, Inches(1.35), Inches(6.4), Inches(1.15))
    write_paragraph(
        tf.paragraphs[0],
        "La app del local para dejar evidencia de cada pedido, leer el código solo y defender los reclamos de PedidosYa, Rappi, Rappi Turbo y Mercado Pago.",
        16,
        False,
        TEXT,
        space_after=0,
    )
    box, tf = textbox(s, MARGIN, Inches(2.6), Inches(6.4), Inches(3.4))
    bullets(
        tf,
        [
            "Dos fotos en el mostrador: ticket y bolsa. El código no se tipea.",
            "Cada reclamo se cruza con la foto y se prepara para refutar en un toque.",
            "Tablero de pedidos, quejas, AWT y plata en disputa, recuperada o perdida.",
            "Ya está en uso en La Plata. El mismo flujo se puede abrir en cualquier local.",
        ],
        15,
    )
    picture(s, ASSETS / "05-pedidos.png", Inches(7.15), Inches(0.55), Inches(5.7), Inches(6.35))
    footer(s, right="delivery.star-app.com.ar  ·  v1.23  ·  En operación en el local")

    # 2 Why
    s = blank_slide(prs)
    kicker(s, "Por qué existe")
    title(s, "Menos fricción en el turno. Más plata que se puede recuperar.")
    lead(s, "Hoy el agregador descuenta el reclamo si el local no tiene evidencia a tiempo. La app cierra ese hueco.")
    stats = [
        ("857", "pedidos · semana en curso", HEAD),
        ("$ 63.096", "en quejas esta semana · aún sin recuperar", WARN),
        ("$ 998.798", "en reclamos del mes · historial del local", HEAD),
        ("5,13%", "AWT vs objetivo 11%", OK),
    ]
    for i, (val, lab, col) in enumerate(stats):
        stat_card(s, MARGIN + Inches(i * 3.15), Inches(1.9), Inches(3.0), Inches(1.2), val, lab, col)

    card(s, MARGIN, Inches(3.3), Inches(6.1), Inches(3.55))
    box, tf = textbox(s, Inches(0.68), Inches(3.45), Inches(5.7), Inches(0.28))
    write_paragraph(tf.paragraphs[0], "OPERACIÓN", 11, True, RED, space_after=0)
    box, tf = textbox(s, Inches(0.68), Inches(3.72), Inches(5.7), Inches(0.36))
    write_paragraph(tf.paragraphs[0], "El mostrador no se frena", 18, True, HEAD, space_after=0)
    box, tf = textbox(s, Inches(0.68), Inches(4.15), Inches(5.7), Inches(2.5))
    bullets(
        tf,
        [
            "El chico saca dos fotos y sigue. La lectura corre atrás.",
            "Si cierran o actualizan la app, la cola se termina igual.",
            "Queda quién sacó la foto. No hay planillas ni códigos mal copiados.",
            "A las 13:30 el listado de quejas se cruza solo, en todos los dispositivos.",
        ],
        14,
    )

    card(s, Inches(6.85), Inches(3.3), Inches(6.0), Inches(3.55))
    box, tf = textbox(s, Inches(7.05), Inches(3.45), Inches(5.6), Inches(0.28))
    write_paragraph(tf.paragraphs[0], "ECONOMÍA", 11, True, OK, space_after=0)
    box, tf = textbox(s, Inches(7.05), Inches(3.72), Inches(5.6), Inches(0.36))
    write_paragraph(tf.paragraphs[0], "Cada foto es una chance de no perder el descuento", 16, True, HEAD, space_after=0)
    box, tf = textbox(s, Inches(7.05), Inches(4.15), Inches(5.6), Inches(2.5))
    bullets(
        tf,
        [
            "Sin foto, el reclamo se paga. Con foto, se puede refutar.",
            "Un toque descarga la evidencia, copia el código y abre el portal.",
            "Se ve la plata en disputa, la recuperada (Ref. aceptado) y la perdida.",
            "% de quejas y AWT contra objetivo, por agregador y por período.",
        ],
        14,
    )
    footer(s)

    # 3 Flow
    s = blank_slide(prs)
    kicker(s, "Flujo")
    title(s, "Cinco pasos. De la bolsa al tablero.")
    lead(s, "No hay un sistema paralelo. Es el mismo recorrido todos los días, en el celular o en la PC.")
    steps = [
        ("01 · Captura", "Ticket", "El recuadro pide el renglón de CODIGO. Flash si hace falta. Nombre de quien saca la foto."),
        ("02 · Captura", "Evidencia", "Bolsa, contenido y ticket a la vista. Esa foto es la defensa si llega el reclamo."),
        ("03 · Lectura", "Código solo", "OCR en segundo plano: PEYA, RAPPI, RAPPITURBO, MPD. Si el celular se cierra, sigue al volver."),
        ("04 · Reclamos", "Cruce", "Se pega o se lee el Sheet. Cada queja se une a la foto. Queda en el historial."),
        ("05 · Plata", "Refutar y medir", "Preparar para refutar. Estados Queja → Refutado → Aceptado o Rechazado."),
    ]
    for i, (n, h3, body) in enumerate(steps):
        x = MARGIN + Inches(i * 2.5)
        card(s, x, Inches(1.9), Inches(2.38), Inches(2.15))
        box, tf = textbox(s, x + Inches(0.12), Inches(2.0), Inches(2.14), Inches(0.28))
        write_paragraph(tf.paragraphs[0], n.upper(), 10, True, RED, space_after=0)
        box, tf = textbox(s, x + Inches(0.12), Inches(2.28), Inches(2.14), Inches(0.32))
        write_paragraph(tf.paragraphs[0], h3, 16, True, HEAD, space_after=0)
        box, tf = textbox(s, x + Inches(0.12), Inches(2.62), Inches(2.14), Inches(1.3))
        write_paragraph(tf.paragraphs[0], body, 12, False, TEXT, space_after=0)
    picture(s, ASSETS / "01-metricas.png", MARGIN, Inches(4.2), Inches(12.35), Inches(2.8))
    footer(s)

    # 4 Phone capture
    s = blank_slide(prs)
    kicker(s, "Mostrador")
    title(s, "Pensada para el celular, en el momento de armar el pedido.")
    picture(s, ASSETS / "02-captura-mobile.png", Inches(0.55), Inches(1.35), Inches(2.55), Inches(5.5))
    picture(s, ASSETS / "10-camara-mobile.png", Inches(3.3), Inches(1.35), Inches(2.55), Inches(5.5))
    box, tf = textbox(s, Inches(6.15), Inches(1.5), Inches(6.7), Inches(5.2))
    bullets(
        tf,
        [
            ("Quién saca la foto", " primero. Queda registrado en cada pedido."),
            ("Sacar foto", " grande y centrado. Cámara en dos pasos: ticket y evidencia, sin cerrarse entre pedidos."),
            ("Flash", " abajo, fácil de tocar. Guía para encuadrar el código."),
            ("Cola de subida: ", "pueden seguir armando. Si cierran la app, los pedidos en cola se terminan de leer."),
            "También se puede elegir de la galería o cargar un archivo (remito, Excel) desde la PC.",
        ],
        16,
    )
    footer(s)

    # 5 OCR
    s = blank_slide(prs)
    kicker(s, "Lectura automática")
    title(s, "El código se lee del ticket. Nadie lo copia a mano.")
    lead(s, "Menos error, menos “código no encontrado”, y el archivo ya nace con el número de pedido.")
    shots = [
        (ASSETS / "ejemplo-peya.jpg", "PedidosYa", "CODIGO: PEYA-2286878556", "Pedido #PEYA-2286878556"),
        (ASSETS / "ejemplo-rappi.jpg", "Rappi", "CODIGO: RAPPI-480403041", "Pedido #RAPPI-480403041"),
        (ASSETS / "ejemplo-bolsa.jpg", "Evidencia", "Bolsa, producto y ticket", "Si el ticket no se lee, la foto se guarda igual"),
    ]
    for i, (img, h3, line, code) in enumerate(shots):
        x = MARGIN + Inches(i * 4.2)
        picture(s, img, x, Inches(1.9), Inches(4.0), Inches(3.7))
        box, tf = textbox(s, x, Inches(5.7), Inches(4.0), Inches(0.32))
        write_paragraph(tf.paragraphs[0], h3, 16, True, HEAD, space_after=0)
        box, tf = textbox(s, x, Inches(6.02), Inches(4.0), Inches(0.28))
        write_paragraph(tf.paragraphs[0], line, 12, False, MUTED, space_after=0)
        box, tf = textbox(s, x, Inches(6.3), Inches(4.0), Inches(0.32))
        write_paragraph(tf.paragraphs[0], code, 13, True, OK, space_after=0)
    footer(s)

    # 6 Gallery
    s = blank_slide(prs)
    kicker(s, "Galería")
    title(s, "Cada pedido queda identificado y a un toque.")
    picture(s, ASSETS / "05-pedidos.png", MARGIN, Inches(1.35), Inches(6.15), Inches(5.45))
    picture(s, ASSETS / "03-lightbox.png", Inches(6.85), Inches(1.35), Inches(6.0), Inches(5.45))
    box, tf = textbox(s, MARGIN, Inches(6.88), Inches(8.2), Inches(0.28))
    write_paragraph(tf.paragraphs[0], "PedidosYa, Rappi, Rappi Turbo y Mercado Pago  ·  En vivo  ·  ver, editar, descargar", 12, False, MUTED, space_after=0)
    box, tf = textbox(s, Inches(8.4), Inches(6.88), Inches(4.4), Inches(0.28))
    write_paragraph(tf.paragraphs[0], "Desde la galería se puede marcar un pedido como reclamo", 12, False, MUTED, PP_ALIGN.RIGHT, 0)
    footer(s)

    # 7 Complaints
    s = blank_slide(prs)
    kicker(s, "Reclamos")
    title(s, "Del listado del agregador a “Preparar para refutar”.")
    box, tf = textbox(s, MARGIN, Inches(1.35), Inches(5.7), Inches(5.0))
    bullets(
        tf,
        [
            ("Importar", " el Sheet o pegar celdas (código, hora, monto, combo, motivo)."),
            ("Cruzar", " con las fotos. A las 13:30 se hace solo, una vez, para todos los celulares del local."),
            ("Historial: ", "Queja, Refutado, Ref. aceptado o Ref. rechazado. Filtro por día, agregador y estado."),
            ("Preparar para refutar: ", "descarga la foto, copia el código y abre PedidosYa Portal o Rappi Partners."),
            "Si el pedido se marcó como reclamo desde la galería, ya entra al historial aunque todavía falten los datos del Sheet.",
        ],
        15,
    )
    picture(s, ASSETS / "07-historial.png", Inches(6.4), Inches(1.3), Inches(6.45), Inches(5.4))
    box, tf = textbox(s, MARGIN, Inches(6.55), Inches(8.0), Inches(0.3))
    write_paragraph(tf.paragraphs[0], "Mes en curso  ·  78 reclamos  ·  $ 998.798 en el historial", 12, True, HEAD, space_after=0)
    box, tf = textbox(s, Inches(8.3), Inches(6.55), Inches(4.5), Inches(0.3))
    write_paragraph(tf.paragraphs[0], "Sin foto = plata que hoy no se puede defender", 12, False, MUTED, PP_ALIGN.RIGHT, 0)
    footer(s)

    # 8 Metrics
    s = blank_slide(prs)
    kicker(s, "Gerencia")
    title(s, "Métricas: operación y montos, por agregador.")
    metrics = [
        ("2,45%", "quejas · objetivo 2,40%", HEAD),
        ("5,13%", "AWT · objetivo 11%", OK),
        ("21", "quejas en la semana", HEAD),
        ("$ 0", "recuperado · oportunidad de refutar", HEAD),
    ]
    for i, (val, lab, col) in enumerate(metrics):
        stat_card(s, MARGIN + Inches(i * 3.15), Inches(1.35), Inches(3.0), Inches(1.15), val, lab, col)
    picture(s, ASSETS / "01-metricas.png", MARGIN, Inches(2.7), Inches(12.35), Inches(4.15))
    footer(s, right="Hoy, Ayer, 7 días, Semana, Mes o rango  ·  corte por agregador")

    # 9 Expansion
    s = blank_slide(prs)
    kicker(s, "Red de locales")
    title(s, "De La Plata a todos los locales, sin un sistema nuevo.")
    card(s, MARGIN, Inches(1.35), Inches(6.1), Inches(2.85))
    box, tf = textbox(s, Inches(0.68), Inches(1.48), Inches(5.7), Inches(0.26))
    write_paragraph(tf.paragraphs[0], "QUÉ HAY QUE HACER", 11, True, RED, space_after=0)
    box, tf = textbox(s, Inches(0.68), Inches(1.74), Inches(5.7), Inches(0.32))
    write_paragraph(tf.paragraphs[0], "El mismo link, el mismo flujo", 16, True, HEAD, space_after=0)
    box, tf = textbox(s, Inches(0.68), Inches(2.1), Inches(5.7), Inches(1.95))
    bullets(
        tf,
        [
            "Se abre delivery.star-app.com.ar en el celular del local y se instala en el inicio. No hay usuarios ni contraseñas.",
            "El equipo saca las mismas dos fotos. No hay capacitación extra ni otro software.",
            "El Sheet de reclamos ya trae la columna de local: el cruce y el historial siguen el mismo método.",
            "Un local nuevo puede empezar el mismo día: ícono en el teléfono y cámara.",
        ],
        12,
    )

    card(s, Inches(6.85), Inches(1.35), Inches(6.0), Inches(2.85))
    box, tf = textbox(s, Inches(7.05), Inches(1.48), Inches(5.6), Inches(0.26))
    write_paragraph(tf.paragraphs[0], "QUÉ GANA LA EMPRESA", 11, True, OK, space_after=0)
    box, tf = textbox(s, Inches(7.05), Inches(1.74), Inches(5.6), Inches(0.32))
    write_paragraph(tf.paragraphs[0], "Un estándar, no una isla", 16, True, HEAD, space_after=0)
    box, tf = textbox(s, Inches(7.05), Inches(2.1), Inches(5.6), Inches(1.95))
    bullets(
        tf,
        [
            "Misma evidencia, mismos estados, mismas métricas en cada punto de venta.",
            "La plata en disputa se ve. La recuperada también. Se puede comparar local contra local.",
            "Menos pedidos “sin foto” en la red = menos descuentos que se dan por no poder refutar.",
            "No se paga una app distinta por sucursal. Se replica lo que ya funciona en La Plata.",
        ],
        12,
    )
    picture(s, ASSETS / "11-galeria-mobile.png", Inches(1.55), Inches(4.35), Inches(2.35), Inches(2.55))
    picture(s, ASSETS / "12-historial-mobile.png", Inches(5.5), Inches(4.35), Inches(2.35), Inches(2.55))
    picture(s, ASSETS / "08-metricas-mobile.png", Inches(9.4), Inches(4.35), Inches(2.35), Inches(2.55))
    footer(s)

    # 10 Close
    s = blank_slide(prs)
    kicker(s, "Siguiente paso")
    title(s, "Ya está en la calle. El siguiente local es instalarla.")
    box, tf = textbox(s, MARGIN, Inches(1.35), Inches(6.2), Inches(1.1))
    write_paragraph(
        tf.paragraphs[0],
        "La Plata ya registra pedidos, lee códigos, cruza reclamos y mide quejas y AWT. Expandirlo es copiar ese hábito, no armar otro proyecto.",
        16,
        False,
        TEXT,
        space_after=0,
    )
    box, tf = textbox(s, MARGIN, Inches(2.5), Inches(6.2), Inches(2.4))
    bullets(
        tf,
        [
            "App web: celular y escritorio 16:9. Se instala como aplicación.",
            "Cuatro agregadores. Cámara, galería, reclamos, historial, métricas y archivos.",
            "Pedido con foto = chance de recuperar el descuento. Pedido sin foto = se paga.",
            "Un link para todos los locales.",
        ],
        15,
    )
    box, tf = textbox(s, MARGIN, Inches(5.15), Inches(6.2), Inches(0.4))
    write_paragraph(tf.paragraphs[0], "delivery.star-app.com.ar", 22, True, HEAD, space_after=0)

    cta = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, Inches(5.65), Inches(2.7), Inches(0.48))
    fill_shape(cta, RED)
    try:
        cta.adjustments[0] = 0.15
    except Exception:
        pass
    tf = cta.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(run, "Abrir la aplicación", 14, True, WHITE)
    tf._txBody.bodyPr.set("anchor", "ctr")

    box, tf = textbox(s, MARGIN, Inches(6.25), Inches(6.2), Inches(0.4))
    write_paragraph(tf.paragraphs[0], "Archivo editable: Delivery-La-Plata.pptx", 13, False, MUTED, space_after=0)

    picture(s, ASSETS / "03-lightbox.png", Inches(6.85), Inches(1.3), Inches(6.0), Inches(5.5))
    footer(s)

    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
